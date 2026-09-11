from __future__ import annotations

import hashlib
import logging
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from .models import FileInventoryItem, SourceInventory
from .storage import sha256_file

logging.getLogger("pypdf").setLevel(logging.ERROR)


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".css",
    ".html",
    ".sql",
    ".r",
    ".m",
    ".sh",
    ".ps1",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
SUPPORTED = {".docx", ".xlsx", ".pptx", ".pdf", *IMAGE_EXTENSIONS, *TEXT_EXTENSIONS}


def _material_type(path: Path) -> str | None:
    name = path.name.lower()
    for token, label in (
        ("实验", "实验方法与数据处理参考"),
        ("课件", "课件"),
        ("ppt", "课件"),
        ("题", "习题与练习"),
        ("复习", "复习资料"),
        ("笔记", "学习笔记"),
        ("代码", "代码参考"),
    ):
        if token in name:
            return label
    if path.suffix.lower() in {".py", ".js", ".java", ".c", ".cpp", ".m", ".r", ".sql"}:
        return "代码参考"
    if path.suffix.lower() == ".pptx":
        return "课件"
    return None


def _docx_stats(path: Path) -> dict[str, Any]:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    images = len(doc.inline_shapes)
    cells = sum(
        1
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    )
    return {
        "paragraphs": len(doc.paragraphs),
        "characters": len(text.strip()),
        "tables": len(doc.tables),
        "non_empty_cells": cells,
        "images": images,
    }


def _xlsx_stats(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    non_empty = 0
    formulas = 0
    for sheet in workbook.worksheets:
        if sheet.max_row * sheet.max_column > 500_000:
            workbook.close()
            raise ValueError("Spreadsheet scan limit; split workbook before processing")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    non_empty += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
    result = {
        "sheets": len(workbook.sheetnames),
        "sheet_names": workbook.sheetnames,
        "non_empty_cells": non_empty,
        "formulas": formulas,
    }
    workbook.close()
    return result


def _pptx_stats(path: Path) -> dict[str, Any]:
    presentation = Presentation(path)
    chars = 0
    images = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chars += len(shape.text.strip())
            if getattr(shape, "shape_type", None) == 13:
                images += 1
    return {"slides": len(presentation.slides), "characters": chars, "images": images}


def _pdf_stats(path: Path) -> dict[str, Any]:
    reader = PdfReader(path)
    characters = 0
    images = 0
    image_scan_errors = 0
    for page in reader.pages:
        characters += len(page.extract_text() or "")
        try:
            images += len(page.images)
        except Exception:  # noqa: BLE001 -- third-party parser boundary, coverage recorded
            image_scan_errors += 1
    return {
        "pages": len(reader.pages),
        "characters": characters,
        "images": images,
        "image_scan_errors": image_scan_errors,
    }


def _image_stats(path: Path) -> dict[str, Any]:
    with Image.open(path) as image:
        image.verify()
    with Image.open(path) as image:
        return {
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
            "format": image.format,
        }


def _text_stats(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    encoding = "utf-8"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        encoding = "gb18030"
        text = raw.decode(encoding)
    return {
        "encoding": encoding,
        "characters": len(text),
        "non_empty_lines": sum(bool(line.strip()) for line in text.splitlines()),
    }


def _inspect(path: Path) -> tuple[str, dict[str, Any], str | None]:
    extension = path.suffix.lower()
    try:
        if path.stat().st_size > 200_000_000:
            raise ValueError("FILE_SIZE_LIMIT")
        if extension in {".docx", ".xlsx", ".pptx"}:
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 500_000_000:
                    raise ValueError("OFFICE_EXPANSION_LIMIT")
        if extension == ".docx":
            stats = _docx_stats(path)
        elif extension == ".xlsx":
            stats = _xlsx_stats(path)
        elif extension == ".pptx":
            stats = _pptx_stats(path)
        elif extension == ".pdf":
            stats = _pdf_stats(path)
        elif extension in IMAGE_EXTENSIONS:
            stats = _image_stats(path)
        elif extension in TEXT_EXTENSIONS:
            stats = _text_stats(path)
        else:
            return "UNSUPPORTED", {}, None
        return "READABLE", stats, None
    except Exception as exc:  # noqa: BLE001 -- a corrupt file must not stop the package
        return "FAILED", {}, type(exc).__name__


def scan_materials(root: Path) -> SourceInventory:
    items: list[FileInventoryItem] = []
    aggregate = hashlib.sha256()
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(root).as_posix()
        try:
            digest = sha256_file(path)
            size = path.stat().st_size
            status, stats, error = _inspect(path)
        except Exception as exc:  # noqa: BLE001 -- capture filesystem/parser failures per file
            digest, size, status, stats, error = "", 0, "FAILED", {}, type(exc).__name__
        aggregate.update(relative.encode("utf-8"))
        aggregate.update(digest.encode("ascii"))
        items.append(
            FileInventoryItem(
                relative_path=relative,
                sha256=digest,
                size_bytes=size,
                extension=path.suffix.lower() or "[none]",
                status=status,
                error=error,
                stats=stats,
                detected_material_type=_material_type(path),
            )
        )
    readable = sum(item.status == "READABLE" for item in items)
    unsupported = sum(item.status == "UNSUPPORTED" for item in items)
    failed = sum(item.status == "FAILED" for item in items)
    warnings = []
    if unsupported:
        warnings.append(
            f"{unsupported} 个文件格式暂不支持深度解析，但已记录哈希和大小。"
        )
    if failed:
        warnings.append(f"{failed} 个文件无法打开；其他文件仍已继续处理。")
    return SourceInventory(
        root_sha256=aggregate.hexdigest(),
        total_files=len(items),
        readable_files=readable,
        unsupported_files=unsupported,
        failed_files=failed,
        items=items,
        warnings=warnings,
    )
