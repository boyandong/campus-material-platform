from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .models import Provenance, QuestionnaireField, QuestionnaireReport
from .storage import sha256_file

FIELD_DEFINITIONS: dict[str, tuple[str, tuple[str, ...]]] = {
    "contact_method": (
        "联系方式类型",
        ("联系方式类型", "卖家联系方式", "联系渠道", "微信/QQ/手机号"),
    ),
    "contact_account": ("联系方式", ("账号或号码", "联系账号", "账号", "号码")),
    "submitted_date": ("填写日期", ("填写日期", "日期")),
    "course_name": ("课程名称", ("课程名称", "课程名")),
    "college": ("学院", ("来源学院", "学院")),
    "major": ("专业", ("来源专业", "专业")),
    "year": ("年份", ("资料年份", "年份", "学年")),
    "teacher": ("任课教师", ("任课教师", "任课老师", "教师", "老师")),
    "material_types": (
        "资料类型",
        ("资料类型", "包含资料类型", "资料形式", "具体包含哪些内容"),
    ),
    "content_summary": ("内容概述", ("内容概述", "资料内容", "具体内容")),
    "completeness": ("完整情况", ("完整情况", "是否完整", "完整度", "完整性")),
    "claimed_file_count": (
        "文件数量",
        ("文件数量", "大概文件数量", "文件数", "文件／份数", "文件/份数"),
    ),
    "chapters_or_experiments": (
        "章节或实验",
        ("包含的章节、实验项目或主要文件", "章节或实验", "章节", "实验项目"),
    ),
    "missing_status": (
        "缺失情况",
        ("是否缺少任何章节、实验、附件或文件", "缺失情况", "是否缺失", "缺少内容"),
    ),
    "missing_details": (
        "缺失说明",
        ("如有缺少，请具体说明", "缺失说明", "缺失详情", "具体缺失"),
    ),
    "historical_grade": (
        "历史成绩或评价",
        ("原课程中的成绩／评价", "历史成绩", "课程成绩", "成绩或评价", "得分"),
    ),
    "grade_evidence": (
        "成绩证明",
        ("可提供成绩或评价截图供核验", "成绩证明", "是否有成绩截图", "证明截图"),
    ),
    "grade_scope": (
        "成绩适用范围",
        ("对应资料范围或补充说明", "成绩适用范围", "成绩对应范围"),
    ),
    "similar_scope": (
        "相近学院或专业",
        ("已知相近学院／专业", "相近学院", "相近专业", "适用学院", "适用专业"),
    ),
    "compatibility_basis": (
        "适配依据",
        ("经验依据或补充说明", "适配依据", "适配原因", "兼容依据"),
    ),
    "buyer_warnings": (
        "买家注意事项",
        ("是否有需要提前提醒买家的特殊情况", "买家注意事项", "注意事项", "风险提示"),
    ),
    "preview_authorization": (
        "预览授权",
        ("是否授权平台制作并展示", "预览授权", "是否授权预览", "展示授权"),
    ),
    "preview_scope": (
        "预览范围",
        ("如需限制展示范围，请说明", "预览范围", "可展示范围", "授权范围"),
    ),
    "minimum_take_home": (
        "最低到手价",
        ("最低实际到手价", "最低到手价", "最低价格", "价格底线"),
    ),
    "other_courses": (
        "其他课程",
        ("如有，请列出课程名称", "其他课程", "其他可提供课程"),
    ),
    "internal_notes": ("其他备注", ("其他说明（可选）", "其他备注", "备注")),
}

CHECKED = {"☒", "☑", "■", "●", "✓", "✔"}
UNCHECKED = {"□", "☐", "○"}


def _document_lines(path: Path) -> list[str]:
    document = Document(path)
    lines = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append("：".join(cells))
    return lines


def _normalize(text: str) -> str:
    return re.sub(r"\s+", "", text).strip("：:？?、")


def _extract_inline_answer(line: str, alias: str) -> str | None:
    match = re.search(re.escape(alias) + r"\s*[：:]\s*(.+)$", line)
    if match:
        candidate = match.group(1).strip()
        return candidate or None
    return None


def _extract_checkboxes(line: str) -> list[str]:
    choices: list[str] = []
    for match in re.finditer(r"([☒☑■●✓✔□☐○])\s*([^☒☑■●✓✔□☐○]+)", line):
        symbol, label = match.groups()
        if symbol in CHECKED:
            choices.append(label.strip(" ，,、;；"))
    return choices


def _looks_like_question(line: str) -> bool:
    normalized = _normalize(line)
    return any(
        _normalize(alias) in normalized
        for _, aliases in FIELD_DEFINITIONS.values()
        for alias in aliases
    )


def _is_blank_template_value(value: str) -> bool:
    stripped = value.strip()
    if not stripped:
        return True
    if any(symbol in stripped for symbol in UNCHECKED) and not any(
        symbol in stripped for symbol in CHECKED
    ):
        return True
    without_placeholders = re.sub(r"[_＿—\-\s年月日人民币¥￥（）()]", "", stripped)
    if not without_placeholders or without_placeholders in {"元", ":", "："}:
        return True
    if stripped.endswith(("？", "?", "：", ":")):
        return True
    if re.match(r"^[一二三四五六七八九十]+、", stripped):
        return True
    return stripped.startswith(("用途说明", "例如", "感谢填写"))


def _is_answer_label(line: str, alias: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    return compact.startswith(alias) and ("：" in compact or ":" in compact)


def parse_questionnaire(path: Path) -> QuestionnaireReport:
    lines = _document_lines(path)
    fields: dict[str, QuestionnaireField] = {}
    warnings: list[str] = []
    for key, (label, aliases) in FIELD_DEFINITIONS.items():
        found_value: str | list[str] | None = None
        raw_value: str | None = None
        source_ref: str | None = None
        confidence = 0.0
        # Prefer actual labelled answers over earlier explanatory questions mentioning the same word.
        candidates = sorted(
            enumerate(lines),
            key=lambda row: (
                not any(
                    _extract_inline_answer(row[1], alias)
                    and not _is_blank_template_value(
                        _extract_inline_answer(row[1], alias)
                    )
                    for alias in aliases
                )
            ),
        )
        for index, line in candidates:
            normalized_line = _normalize(line)
            matching_alias = next(
                (alias for alias in aliases if _normalize(alias) in normalized_line),
                None,
            )
            if not matching_alias:
                continue
            inline = _extract_inline_answer(line, matching_alias)
            checked = _extract_checkboxes(line)
            if checked:
                found_value = checked
                raw_value = line
                confidence = 0.9
            elif inline and not _is_blank_template_value(inline):
                found_value = inline
                raw_value = inline
                confidence = 0.92
            else:
                window_checked: list[str] = []
                window_raw: str | None = None
                for candidate_line in lines[index + 1 : index + 4]:
                    if _looks_like_question(candidate_line) and (
                        not any(
                            symbol in candidate_line for symbol in CHECKED | UNCHECKED
                        )
                        or re.search(r"[:：].*[☒☑□☐]", candidate_line)
                    ):
                        break
                    candidate_checked = _extract_checkboxes(candidate_line)
                    if candidate_checked:
                        window_checked.extend(candidate_checked)
                        window_raw = candidate_line
                if window_checked:
                    found_value = window_checked
                    raw_value = window_raw
                    confidence = 0.78
            if (
                found_value is None
                and _is_answer_label(line, matching_alias)
                and index + 1 < len(lines)
                and not _looks_like_question(lines[index + 1])
                and not _is_blank_template_value(lines[index + 1])
            ):
                next_line = lines[index + 1].strip()
                found_value = next_line
                raw_value = next_line
                confidence = 0.72
            source_ref = f"questionnaire.docx#paragraph-{index + 1}"
            if found_value:
                break
        unknown_values = {
            "",
            "不清楚",
            "不确定",
            "未知",
            "unknown",
            "无法确认",
            "目前不确定",
        }
        selected_values = (
            found_value if isinstance(found_value, list) else [found_value]
        )
        unknown = found_value is None or all(
            str(value or "").strip().lower() in unknown_values
            for value in selected_values
        )
        fields[key] = QuestionnaireField(
            key=key,
            label=label,
            value=None if unknown else found_value,
            raw_value=raw_value,
            provenance=Provenance.UNKNOWN if unknown else Provenance.SELLER_DECLARED,
            source_refs=[source_ref] if source_ref else [],
            confidence=0.0 if unknown else confidence,
        )
    unknown_fields = [
        key for key, field in fields.items() if field.provenance == Provenance.UNKNOWN
    ]
    if unknown_fields:
        warnings.append(f"{len(unknown_fields)} 个字段未能确定，已保留为 UNKNOWN。")
    return QuestionnaireReport(
        source_sha256=sha256_file(path),
        fields=fields,
        unknown_fields=unknown_fields,
        parser_warnings=warnings,
    )
