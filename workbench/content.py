"""Local content extraction; text stays in memory until sanitized for a report/model."""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader

XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, recover=False)
CODE_EXTENSIONS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".sql",
    ".r",
    ".m",
    ".sh",
    ".ps1",
}


@dataclass
class Content:
    text: str = ""
    headings: list[str] = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def extract_content(path: Path) -> Content:
    result = Content()
    ext = path.suffix.lower()
    if ext in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            texts, cell_texts, paragraph_texts = [], [], []
            result.stats["embedded_images"] = sum("/media/" in name for name in names)
            result.stats["embedded_objects"] = sum(
                "/embeddings/" in name for name in names
            )
            result.stats["comments"] = 0
            result.stats["revisions"] = 0
            for name in names:
                if not name.endswith(".xml") or name.startswith("docProps/"):
                    continue
                xml = etree.fromstring(archive.read(name), XML_PARSER)
                if "comments" in name.lower():
                    result.stats["comments"] += len(xml)
                result.stats["revisions"] += len(
                    xml.xpath('//*[local-name()="del" or local-name()="ins"]')
                )
                paragraphs = xml.xpath('//*[local-name()="p"]')
                for paragraph in paragraphs:
                    text = "".join(
                        paragraph.xpath(
                            './/*[local-name()="t" or local-name()="delText"]/text()'
                        )
                    )
                    if text:
                        texts.append(text)
                        paragraph_texts.append(text)
                if not paragraphs:
                    texts.extend(
                        xml.xpath('//*[local-name()="t" or local-name()="v"]/text()')
                    )
                for cell in xml.xpath('//*[local-name()="tc"]'):
                    cell_texts.append(
                        "".join(cell.xpath('.//*[local-name()="t"]/text()'))
                    )
                result.stats["formulas"] = result.stats.get("formulas", 0) + len(
                    xml.xpath('//*[local-name()="oMath" or local-name()="f"]')
                )
            result.text = "\n".join(texts)
            result.stats.update(
                characters=len(result.text.strip()),
                cells=len(cell_texts),
                non_empty_cells=sum(bool(x.strip()) for x in cell_texts),
                substantial_paragraphs=sum(
                    len(x.strip()) >= 40 for x in paragraph_texts
                ),
            )
            for name in ("docProps/core.xml", "docProps/custom.xml"):
                if name in names:
                    xml = etree.fromstring(archive.read(name), XML_PARSER)
                    for element in xml.iter():
                        if element.text and element.text.strip():
                            result.metadata[etree.QName(element).localname] = (
                                element.text
                            )
            if result.stats["embedded_objects"]:
                result.warnings.append("EMBEDDED_OBJECT_REQUIRES_REVIEW")
        if ext == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False)
            text_rows, non_empty, numeric, populated_rows, cells = [], 0, 0, 0, 0
            for sheet in book:
                if sheet.max_row * sheet.max_column > 500_000:
                    result.warnings.append("SHEET_SCAN_LIMIT")
                    continue
                for row in sheet.iter_rows():
                    values = [cell.value for cell in row]
                    cells += len(values)
                    if any(value not in (None, "") for value in values):
                        populated_rows += 1
                        text_rows.append(
                            " | ".join("" if v is None else str(v) for v in values)
                        )
                    non_empty += sum(v not in (None, "") for v in values)
                    numeric += sum(isinstance(v, (int, float)) for v in values)
            result.stats.update(
                cells=cells,
                non_empty_cells=non_empty,
                numeric_cells=numeric,
                populated_rows=populated_rows,
                sheets=len(book.sheetnames),
            )
            result.text = "\n".join(text_rows)
            book.close()
    elif ext == ".pdf":
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        result.text = "\n".join(pages)
        result.stats = {
            "pages": len(pages),
            "page_characters": [len(t.strip()) for t in pages],
        }
        result.metadata = {str(k): str(v) for k, v in (reader.metadata or {}).items()}
    elif ext not in {".png", ".jpg", ".jpeg"}:
        raw = path.read_bytes()
        try:
            result.text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            result.text = raw.decode("gb18030")
    result.stats["characters"] = len(result.text.strip())
    result.stats["placeholder_count"] = len(
        re.findall(r"_{3,}|＿{3,}|待填写|请填写|在此输入|点击此处", result.text)
    )
    result.stats["numeric_tokens"] = len(
        re.findall(r"(?<!\w)\d+(?:\.\d+)?", result.text)
    )
    result.headings = [
        line.strip()[:120]
        for line in result.text.splitlines()
        if re.match(
            r"^\s*(第.{1,8}[章节]|实验\s*[一二三四五六七八九十\d]+|Chapter\s+\d+)",
            line,
            re.IGNORECASE,
        )
    ][:80]
    return result
