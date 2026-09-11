from __future__ import annotations

import io
import json
import stat
import zipfile
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from test_workbench import DisabledProvider, cleanup_read_only

from workbench.buyer_card import build_buyer_card
from workbench.config import Settings
from workbench.content import Content, extract_content
from workbench.database import Database
from workbench.deepseek import DeepSeekProvider
from workbench.duplicates import compare_packages
from workbench.imaging import ImageProcessor
from workbench.models import PackageStatus
from workbench.parser import parse_questionnaire
from workbench.pipeline import WorkbenchPipeline
from workbench.previews import authorized, generate_previews
from workbench.pricing import calculate_pricing
from workbench.privacy import Redactor
from workbench.quality import classify_content
from workbench.redaction import redact_sources
from workbench.scanner import scan_materials
from workbench.storage import extract_zip_read_only, sha256_file


class NoOCR:
    def read(self, image):
        return []


def test_actual_blank_template_and_synthetic_filled_template(tmp_path):
    template = Path(__file__).resolve().parents[1] / "课程学习资料卖家信息采集问卷.docx"
    blank = parse_questionnaire(template)
    assert all(field.value is None for field in blank.fields.values())
    document = Document(template)
    filled = {
        "账号或号码：": "synthetic-account",
        "课程名称：": "数据结构",
        "学院：": "计算机学院",
        "专业：": "软件工程",
        "学年／年份：": "2025",
        "任课老师：": "不确定",
        "主要内容概述：": "原理与数据处理方法",
        "文件／份数：": "3",
        "包含的章节、实验项目或主要文件：": "第一章，第二章",
        "如有缺少，请具体说明：": "第三章",
        "原课程中的成绩／评价：": "90分",
        "对应资料范围或补充说明：": "第一章",
        "已知相近学院／专业：": "电子信息",
        "经验依据或补充说明：": "同课程名称",
        "如需限制展示范围，请说明：": "前2页",
        "其他说明（可选）：": "仅用于测试",
    }
    for p in document.paragraphs:
        if p.text in filled:
            p.text += filled[p.text]
        p.text = (
            p.text.replace("□ QQ", "☒ QQ")
            .replace("□ 代码", "☒ 代码")
            .replace("□ 复习笔记", "☒ 复习笔记")
        )
        p.text = p.text.replace("□ 仅部分内容", "☒ 仅部分内容").replace(
            "□ 有缺少", "☒ 有缺少"
        )
        p.text = p.text.replace("□ 可以", "☒ 可以").replace("□ 授权", "☒ 授权")
        if p.text.startswith("最低实际到手价："):
            p.text = "最低实际到手价：20元"
    path = tmp_path / "filled-template.docx"
    document.save(path)
    report = parse_questionnaire(path)
    expected = {
        "course_name": "数据结构",
        "college": "计算机学院",
        "major": "软件工程",
        "year": "2025",
        "claimed_file_count": "3",
        "historical_grade": "90分",
        "preview_scope": "前2页",
        "compatibility_basis": "同课程名称",
        "minimum_take_home": "20元",
        "internal_notes": "仅用于测试",
    }
    for key, value in expected.items():
        assert report.fields[key].value == value, key
    assert report.fields["preview_authorization"].value == ["授权"]
    assert report.fields["grade_evidence"].value == ["可以"]
    assert report.fields["teacher"].value is None
    assert set(report.fields["material_types"].value) == {"代码", "复习笔记"}


def test_pdf_rasterized_copy_has_no_original_text_and_pptx_embedded_image(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfReader

    originals = tmp_path / "originals"
    originals.mkdir()
    image = Image.new("RGB", (500, 700), "white")
    ImageDraw.Draw(image).text((20, 30), "anonymous synthetic data", fill="black")
    image.save(originals / "scan.pdf", format="PDF")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Inches(1), Inches(1), width=Inches(3))
    deck.save(originals / "slides.pptx")
    inventory = scan_materials(originals)
    contents = {p.name: extract_content(p) for p in originals.iterdir()}
    report = redact_sources(
        originals,
        tmp_path / "processed",
        inventory,
        contents,
        ImageProcessor(NoOCR(), Redactor()),
    )
    assert report["blocked_files"] == 0
    for row in report["items"]:
        output = Path(row["processed_file"])
        assert row["images"]
        if output.suffix == ".pdf":
            assert not (PdfReader(output).pages[0].extract_text() or "").strip()
        else:
            assert len(Presentation(output).slides) == 1


def test_quality_distinguishes_headers_filled_instructions_code(tmp_path):
    book = Workbook()
    book.active.append(["Time", "Voltage", "Current"])
    path = tmp_path / "header.xlsx"
    book.save(path)
    assert classify_content(".xlsx", extract_content(path))[0] == "EMPTY_TEMPLATE"
    book.active.append([1, 2, 3])
    book.save(path)
    assert classify_content(".xlsx", extract_content(path))[0] == "RAW_DATA"
    assert (
        classify_content(".docx", Content(text="实验要求：请完成以下操作"))[0]
        == "ASSIGNMENT_INSTRUCTION"
    )
    assert classify_content(".py", Content(text="print(1)"))[0] == "CODE"
    assert (
        classify_content(
            ".docx", Content(text="待填写\n____\n____", stats={"placeholder_count": 3})
        )[0]
        == "EMPTY_TEMPLATE"
    )


def test_private_strings_removed_without_recording_values():
    source = "姓名：张三 学号：2026123456 电话：13800138000 QQ：12345678 邮箱：demo@example.com"
    redactor = Redactor(["private-contact"])
    redactor.discover(source)
    output = redactor.clean(source + " private-contact")
    for secret in (
        "张三",
        "2026123456",
        "13800138000",
        "12345678",
        "demo@example.com",
        "private-contact",
    ):
        assert secret not in output
        assert secret not in json.dumps(redactor.findings(source))
    assert "[已遮蔽]" in output


def test_ocr_regions_are_blackened_and_cloud_requires_opt_in():
    class FakeOCR:
        def read(self, image):
            return [
                {
                    "text": "学号：2026123456",
                    "box": [[10, 10], [240, 10], [240, 50], [10, 50]],
                    "confidence": 0.99,
                }
            ]

    class Vision:
        calls = 0

        def assess_image(self, image):
            self.calls += 1
            return None, "offline"

    vision = Vision()
    processor = ImageProcessor(FakeOCR(), Redactor(), vision)
    source = Image.new("RGB", (400, 300), "white")
    safe, report, text = processor.process(source)
    assert safe.getpixel((100, 30)) == (0, 0, 0)
    assert source.getpixel((100, 30)) == (255, 255, 255)
    assert report["masked_regions"] == 1 and "2026123456" not in text
    assert vision.calls == 0
    processor.allow_cloud = True
    processor.max_calls = 1
    processor.process(source)
    processor.process(source)
    assert vision.calls == 1


def test_office_copies_strip_metadata_split_names_and_headers(tmp_path):
    originals = tmp_path / "originals"
    originals.mkdir()
    path = originals / "姓名：张三.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("姓名：张")
    p.add_run("三 学号：2026123456")
    doc.sections[0].header.paragraphs[0].text = "联系电话：13800138000"
    doc.core_properties.author = "private-author"
    doc.add_paragraph("第一章 数据结构\n" + "数据结构方法与概念。" * 30)
    doc.save(path)
    before = sha256_file(path)
    content = extract_content(path)
    report = redact_sources(
        originals,
        tmp_path / "processed",
        scan_materials(originals),
        {path.name: content},
        ImageProcessor(NoOCR(), Redactor()),
    )
    row = report["items"][0]
    assert row["status"] == "NEEDS_REVIEW"
    copy = Path(row["processed_file"])
    Document(copy)
    with zipfile.ZipFile(copy) as archive:
        serialized = b"\n".join(
            archive.read(n) for n in archive.namelist() if n.endswith(".xml")
        ).decode("utf-8")
    for value in ("张三", "2026123456", "13800138000", "private-author"):
        assert value not in serialized
    assert sha256_file(path) == before
    assert copy.name == "F0001.docx"


def test_spreadsheet_numeric_student_id_redaction_remains_openable(tmp_path):
    originals = tmp_path / "originals"
    originals.mkdir()
    book = Workbook()
    book.active.append(["学号", 2026123456])
    book.active.append(["姓名", "张三"])
    path = originals / "test.xlsx"
    book.save(path)
    content = extract_content(path)
    # A table row separator is recognized just like a label delimiter.
    report = redact_sources(
        originals,
        tmp_path / "processed",
        scan_materials(originals),
        {path.name: content},
        ImageProcessor(NoOCR(), Redactor()),
    )
    result = load_workbook(report["items"][0]["processed_file"])
    assert str(result.active.cell(1, 2).value) != "2026123456"
    assert result.active.cell(2, 2).value != "张三"
    result.close()


def test_preview_authorization_and_academic_scope_are_enforced(tmp_path):
    image = tmp_path / "input.png"
    Image.new("RGB", (500, 700), "white").save(image)
    candidate = {
        "file_id": "F0001",
        "ordinal": 1,
        "kind": "image",
        "path": str(image),
        "quality": "readable",
        "score": 0.5,
        "sha256": sha256_file(image),
    }
    privacy = {"candidates": [candidate]}
    quality = {"items": [{"file_id": "F0001", "academic_review_required": True}]}
    assert not authorized("不允许")
    assert not authorized("可能可以")
    assert (
        generate_previews(tmp_path / "out", privacy, quality, None, None)["items"] == []
    )
    assert (
        generate_previews(tmp_path / "out", privacy, quality, "允许", "两张")["status"]
        == "NEEDS_SCOPE_CONFIRMATION"
    )
    assert (
        generate_previews(tmp_path / "out", privacy, quality, "允许", None)["items"]
        == []
    )
    report = generate_previews(
        tmp_path / "out", privacy, quality, "允许", None, {"F0001": [1]}
    )
    assert len(report["items"]) == 1
    with Image.open(report["items"][0]["path"]) as preview:
        assert preview.height < 700


def test_pricing_floor_shares_and_unknown():
    quality = {"effective_content_files": 3, "risk_counts": {}, "counts": {}}
    result = calculate_pricing("20元", quality)
    assert result["seller_share"] == "0.65"
    assert result["platform_share_with_referral"] == "0.25"
    assert Decimal(result["recommended_min_sale_price"]) * Decimal(".65") >= 20
    assert result["confidence"] == "LOW_CONFIDENCE"
    assert calculate_pricing("20-30", quality)["status"] == "UNKNOWN"
    with pytest.raises(ValueError):
        calculate_pricing("20", quality, "0.2", "0.3")


@pytest.mark.parametrize(
    "entry",
    ["C:/escape.txt", "file.txt:stream", "CON.txt", "folder/../bad.txt", "foo."],
)
def test_windows_zip_unsafe_names_are_rejected(tmp_path, entry):
    path = tmp_path / "bad.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(entry, "data")
    with pytest.raises(ValueError):
        extract_zip_read_only(path, tmp_path / "out")


def clean_pipeline(tmp_path):
    questionnaire = tmp_path / "questionnaire.docx"
    document = Document()
    for line in [
        "课程名称：数据结构",
        "学院：计算机学院",
        "专业：软件工程",
        "年份：2025",
        "账号或号码：private-user",
        "内容概述：课程原理和方法概述",
        "资料类型：学习笔记",
        "完整情况：部分资料",
        "文件数量：1",
        "缺失说明：无",
        "预览授权：允许",
        "预览范围：仅脱敏部分预览",
        "最低到手价：20元",
    ]:
        document.add_paragraph(line)
    document.save(questionnaire)
    source = tmp_path / "source"
    source.mkdir()
    (source / "数据结构笔记.txt").write_text(
        "第一章 数据结构\n" + "数据结构的基本概念包括线性结构与树形结构。" * 40,
        encoding="utf-8",
    )
    data = tmp_path / "data"
    data.mkdir()
    db = Database(tmp_path / "tasks.sqlite")
    pipeline = WorkbenchPipeline(data, db, DisabledProvider(), ocr=NoOCR())
    pid = pipeline.create_package(questionnaire, source)
    return pipeline, pid, data / pid


def test_full_flow_approval_manual_answers_and_hash_guard(tmp_path):
    pipeline, pid, package = clean_pipeline(tmp_path)
    try:
        product = pipeline.process(pid)
        assert product.quality_gate["approval_enabled"], product.quality_gate
        for name in (
            "quality_report",
            "privacy_report",
            "pricing_report",
            "preview_report",
            "version_comparison",
            "questions",
            "model_calls",
        ):
            assert (package / "reports" / f"{name}.json").exists()
        with pytest.raises(ValueError):
            pipeline.review(pid, PackageStatus.APPROVED, "Reviewer", "checked")
        checks = {
            key: True
            for key in (
                "privacy",
                "previews",
                "authorization",
                "price",
                "description",
                "warnings",
            )
        }
        pipeline.review(
            pid,
            PackageStatus.APPROVED,
            "Reviewer",
            "逐文件核对，未发现敏感信息。",
            confirmations=checks,
        )
        assert json.loads((package / "buyer_card.json").read_text(encoding="utf-8"))[
            "publishable"
        ]
        assert (
            json.loads((package / "product.json").read_text(encoding="utf-8"))[
                "quality_gate"
            ]["status"]
            == "APPROVED"
        )
        pipeline.answer_question(pid, "field:year", "2026", "Reviewer")
        assert not json.loads(
            (package / "buyer_card.json").read_text(encoding="utf-8")
        )["publishable"]
        updated = pipeline.process(pid)
        assert updated.course["year"].value == "2026"
        assert updated.course["year"].provenance.value == "HUMAN_VERIFIED"
        assert updated.course["year"].verified_by == "Reviewer"
        processed = package / updated.privacy["items"][0]["processed_file"]
        processed.write_text("changed", encoding="utf-8")
        with pytest.raises(ValueError, match="产物"):
            pipeline.review(
                pid, PackageStatus.APPROVED, "Reviewer", "checked", confirmations=checks
            )
    finally:
        cleanup_read_only(package)


def test_free_text_public_card_does_not_leak_contacts_or_paths(tmp_path):
    pipeline, pid, package = clean_pipeline(tmp_path)
    try:
        product = pipeline.process(pid)
        product.materials[
            "summary"
        ].value = "联系 private-user，手机号 13800138000，位置 C:\\private\\file.txt"
        product.warnings = ["底价20元，内部用户private-user"]
        serialized = json.dumps(build_buyer_card(product), ensure_ascii=False)
        for secret in ("private-user", "13800138000", "C:\\private", "20元"):
            assert secret not in serialized
    finally:
        cleanup_read_only(package)


def test_new_stages_recover_as_interrupted(tmp_path):
    db_path = tmp_path / "tasks.sqlite"
    database = Database(db_path)
    database.create_package("PRD-T", tmp_path / "PRD-T")
    database.update_package("PRD-T", PackageStatus.PRIVACY_SCAN, "redaction")
    assert Database(db_path).get_package("PRD-T")["status"] == "INTERRUPTED"


def test_http_error_diagnostics_never_echo_secret(tmp_path):
    settings = Settings(
        tmp_path,
        tmp_path / "data",
        tmp_path / "db",
        "dummy-secret",
        "https://api.deepseek.com",
        "deepseek-v4-flash",
    )
    provider = DeepSeekProvider(settings)
    provider._client.close()
    count = []

    def handler(request):
        count.append(1)
        return httpx.Response(401, json={"error": "dummy-secret"})

    provider._client = httpx.Client(
        transport=httpx.MockTransport(handler), base_url="https://api.deepseek.com"
    )
    result, warning = provider.classify({"content_summary": "test"})
    assert result is None and "HTTP 401" in warning and "dummy-secret" not in warning
    assert len(count) == 1
    assert "dummy-secret" not in json.dumps(provider.call_log)
    provider.close()


def test_ocr_failure_and_pdf_page_cap_block_derivatives(tmp_path):
    class BrokenOCR:
        def read(self, image):
            raise RuntimeError("test failure")

    originals = tmp_path / "originals"
    originals.mkdir()
    image = Image.new("RGB", (500, 700), "white")
    image.save(originals / "one.png")
    image.save(
        originals / "two.pdf", format="PDF", save_all=True, append_images=[image]
    )
    inventory = scan_materials(originals)
    contents = {p.name: extract_content(p) for p in originals.iterdir()}
    report = redact_sources(
        originals,
        tmp_path / "out",
        inventory,
        contents,
        ImageProcessor(BrokenOCR(), Redactor()),
        max_pdf_pages=1,
    )
    assert report["blocked_files"] == 2
    assert not report["candidates"]


def test_vision_invalid_coordinates_never_mask_arbitrary_area():
    from workbench.models import VisionAssessment

    class BadVision:
        def assess_image(self, image):
            return VisionAssessment(
                quality="readable",
                content_type="test",
                privacy_regions=[{"kind": "test", "box": [-1, 0, 2, 1]}],
            ), None

    processor = ImageProcessor(NoOCR(), Redactor(), BadVision(), allow_cloud=True)
    result, report, _ = processor.process(Image.new("RGB", (400, 300), "white"))
    assert "INVALID_VISION_REGION" in report["warnings"]
    assert result.getpixel((200, 150)) == (255, 255, 255)
    assert report["manual_review_required"]


def test_duplicate_packages_detected_by_content_not_file_names(tmp_path):
    pipeline, pid, package = clean_pipeline(tmp_path)
    try:
        pipeline.process(pid)
        quality = json.loads(
            (package / "reports/quality_report.json").read_text(encoding="utf-8")
        )
        result = compare_packages(
            package.parent, "PRD-OTHER", {"name": "数据结构", "year": "2025"}, quality
        )
        assert result["matches"][0]["classification"] == "EXACT_DUPLICATE"
        assert result["automatic_merge"] is False
    finally:
        cleanup_read_only(package)


def test_original_tamper_rejected_before_any_model_call(tmp_path):
    pipeline, pid, package = clean_pipeline(tmp_path)
    try:
        path = next((package / "original/materials").iterdir())
        path.chmod(stat.S_IWRITE | stat.S_IREAD)
        path.write_text("changed", encoding="utf-8")
        with pytest.raises(ValueError, match="原件哈希"):
            pipeline.process(pid)
        assert pipeline.database.get_package(pid)["status"] == "FAILED"
    finally:
        cleanup_read_only(package)


def test_manual_price_recalculates_payout_without_changing_seller_floor(tmp_path):
    pipeline, pid, package = clean_pipeline(tmp_path)
    try:
        pipeline.process(pid)
        pipeline.answer_question(pid, "pricing:list_price", "50", "Reviewer")
        product = pipeline.process(pid)
        result = product.pricing["calculated"].value
        assert result["payout_example"]["seller"] == "32.50"
        assert result["payout_example"]["referrer"] == "5.00"
        assert result["payout_example"]["platform_with_referral"] == "12.50"
        assert result["seller_min_take_home"] == "20"
        assert product.pricing["list_price"].provenance.value == "HUMAN_VERIFIED"
    finally:
        cleanup_read_only(package)


def test_unknown_checked_answer_keeps_raw_value(tmp_path):
    doc = Document()
    doc.add_paragraph("完整性：□ 完整的一套 □ 仅部分内容 ☒ 无法确认")
    path = tmp_path / "unknown.docx"
    doc.save(path)
    result = parse_questionnaire(path)
    assert result.fields["completeness"].value is None
    assert "无法确认" in result.fields["completeness"].raw_value
    assert result.fields["completeness"].provenance.value == "UNKNOWN"
