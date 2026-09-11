from __future__ import annotations

import json
import sqlite3
import stat
import zipfile
from pathlib import Path

import httpx
import pytest
from docx import Document
from PIL import Image

from workbench.buyer_card import build_buyer_card
from workbench.config import Settings
from workbench.database import Database
from workbench.deepseek import DeepSeekProvider
from workbench.models import PackageStatus, ProductDocument, Provenance, SourcedValue
from workbench.parser import parse_questionnaire
from workbench.pipeline import WorkbenchPipeline
from workbench.scanner import scan_materials
from workbench.storage import extract_zip_read_only, sha256_file


def make_questionnaire(path: Path, *, contradictory: bool = True) -> None:
    document = Document()
    values = [
        "联系方式类型：微信",
        "账号或号码：synthetic-contact",
        "课程名称：数据结构",
        "学院：计算机学院",
        "专业：软件工程",
        "年份：2025",
        "任课教师：不确定",
        "资料类型：☒ 课件 □ 题库 ☒ 实验资料",
        "内容概述：包含课程课件、实验方法和数据处理思路",
        "完整情况：完整",
        "文件数量：5",
        "缺失情况：有缺失" if contradictory else "缺失情况：无",
        "缺失说明：缺少一次实验" if contradictory else "缺失说明：无",
        "历史成绩：90分",
        "成绩证明：无",
        "相近学院或专业：电子信息学院",
        "适配依据：课程名相同，但教师可能不同",
        "买家注意事项：请先核对章节",
        "预览授权：允许",
        "预览范围：脱敏后两页",
        "最低到手价：20元",
        "其他备注：内部合成样本",
    ]
    for value in values:
        document.add_paragraph(value)
    document.save(path)


class DisabledProvider:
    review_model = "deepseek-v4-pro"
    def classify(self, safe_free_text):
        return None, "未配置 DEEPSEEK_API_KEY，语义归类已进入人工审核。"
    def advanced_review(self, safe_conflicts):
        return None, "高级审核未调用。"


class StubOCR:
    def read(self, image):
        return []


def cleanup_read_only(root: Path) -> None:
    for path in root.rglob("*"):
        if path.is_file():
            path.chmod(stat.S_IWRITE | stat.S_IREAD)


def test_questionnaire_parser_preserves_unknown_and_checkbox_values(tmp_path: Path):
    questionnaire = tmp_path / "questionnaire.docx"
    make_questionnaire(questionnaire)
    report = parse_questionnaire(questionnaire)
    assert report.fields["course_name"].value == "数据结构"
    assert report.fields["teacher"].provenance == Provenance.UNKNOWN
    assert "课件" in report.fields["material_types"].value
    assert "实验资料" in report.fields["material_types"].value
    assert report.fields["contact_account"].value == "synthetic-contact"


def test_blank_template_placeholders_are_not_treated_as_answers(tmp_path: Path):
    questionnaire = tmp_path / "blank.docx"
    document = Document()
    for line in (
        "卖家联系方式（QQ、微信、手机号任选其一）", "□ QQ □ 微信 □ 手机号",
        "账号或号码：", "填写日期：________年____月____日",
        "课程名称：", "学院：", "专业：", "学年／年份：",
        "资料类型：□ 课件 □ 题库", "最低实际到手价：人民币（¥）____________ 元",
    ):
        document.add_paragraph(line)
    document.save(questionnaire)
    report = parse_questionnaire(questionnaire)
    for key in ("contact_method", "contact_account", "submitted_date", "course_name", "college", "major", "year", "material_types", "minimum_take_home"):
        assert report.fields[key].provenance == Provenance.UNKNOWN


def test_zip_path_traversal_is_rejected(tmp_path: Path):
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("../escape.txt", "bad")
    with pytest.raises(ValueError, match="unsafe path"):
        extract_zip_read_only(archive, tmp_path / "output")
    assert not (tmp_path / "escape.txt").exists()


def test_scanner_continues_after_corrupt_and_unsupported_files(tmp_path: Path):
    root = tmp_path / "materials"
    root.mkdir()
    (root / "notes.txt").write_text("data structures", encoding="utf-8")
    (root / "broken.pdf").write_bytes(b"not-a-pdf")
    (root / "legacy.ms13").write_bytes(b"legacy")
    result = scan_materials(root)
    assert result.total_files == 3
    assert result.readable_files == 1
    assert result.failed_files == 1
    assert result.unsupported_files == 1


def test_buyer_card_uses_public_whitelist():
    field = lambda value, provenance=Provenance.SELLER_DECLARED: SourcedValue(value=value, provenance=provenance, confidence=1)
    product = ProductDocument(
        product_id="PRD-TEST",
        seller={"contact_account": field("private-wechat")},
        course={"name": field("数据结构"), "college": field("计算机学院"), "major": field("软件工程"), "year": field("2025"), "teacher": field(None)},
        materials={"types": field(["课件"]), "summary": field("摘要"), "file_count": field(1), "readable_file_count": field(1), "extensions": field({".pdf": 1})},
        quality={"declared_completeness": field("完整"), "unreadable_files": field(0)},
        compatibility={"scope": field("本专业"), "basis": field("课程相同")},
        authorization={"preview_authorization": field("允许")},
        pricing={"minimum_take_home": field("20元")},
        historical_result={"seller_declared": field("90分")},
        warnings=[], unanswered_questions=[],
        quality_gate={"status": "NEEDS_MANUAL_REVIEW", "approval_enabled": False}, artifacts={},
    )
    card = build_buyer_card(product)
    serialized = json.dumps(card, ensure_ascii=False)
    assert "private-wechat" not in serialized
    assert "20元" not in serialized
    assert "卖家声明，未经核验" in serialized
    assert "不构成成绩、提分或通过保证" in serialized


def test_deepseek_failure_retries_once_then_falls_back(tmp_path: Path):
    settings = Settings(tmp_path, tmp_path / "data", tmp_path / "db.sqlite", "test-key", "https://api.deepseek.com", "deepseek-v4-flash")
    provider = DeepSeekProvider(settings)

    class FailingClient:
        calls = 0
        def post(self, *args, **kwargs):
            self.calls += 1
            raise httpx.ReadTimeout("timeout")
        def close(self):
            pass

    failing = FailingClient()
    provider._client = failing
    result, warning = provider.classify({"content_summary": "合成内容"})
    assert result is None
    assert failing.calls == 2
    assert "失败" in warning


def test_deepseek_routes_default_and_advanced_models(tmp_path: Path):
    settings = Settings(
        tmp_path, tmp_path / "data", tmp_path / "db.sqlite", "test-key", "https://api.deepseek.com",
        "deepseek-v4-flash", "deepseek-v4-flash-vision-exp", "deepseek-v4-pro",
    )
    provider = DeepSeekProvider(settings)

    class Response:
        def __init__(self, content):
            self.content = content
        def raise_for_status(self):
            return None
        def json(self):
            return {"choices": [{"message": {"content": json.dumps(self.content, ensure_ascii=False)}}]}

    class RoutingClient:
        def __init__(self):
            self.models = []
        def post(self, path, json):
            self.models.append(json["model"])
            if json["model"] == "deepseek-v4-pro":
                return Response({"risk_level": "HIGH", "conflict_summary": "文件冲突", "recommended_actions": ["人工核对"], "requires_human_review": True, "confidence": 0.8})
            return Response({"normalized_material_types": ["课件"], "material_summary": "课程资料", "catalog_outline": ["第一章"], "pricing_suggestion": "仅作内部参考", "compatibility_summary": None, "buyer_warning_summary": None, "confidence": 0.7})
        def close(self):
            pass

    client = RoutingClient()
    provider._client = client
    semantic, semantic_warning = provider.classify({"content_summary": "合成内容"})
    advanced, advanced_warning = provider.advanced_review({"conflicts": [{"code": "TEST"}]})
    assert semantic_warning is None and semantic.material_summary == "课程资料"
    assert advanced_warning is None and advanced.risk_level == "HIGH"
    assert client.models == ["deepseek-v4-flash", "deepseek-v4-pro"]
    assert provider.vision_model == "deepseek-v4-flash-vision-exp"


def test_database_marks_in_progress_job_interrupted_after_restart(tmp_path: Path):
    db_path = tmp_path / "jobs.sqlite"
    db = Database(db_path)
    package = tmp_path / "PRD-INTERRUPT"
    package.mkdir()
    db.create_package("PRD-INTERRUPT", package)
    db.update_package("PRD-INTERRUPT", PackageStatus.SCANNING_FILES, "scan")
    restarted = Database(db_path)
    assert restarted.get_package("PRD-INTERRUPT")["status"] == PackageStatus.INTERRUPTED.value


def test_end_to_end_pipeline_generates_all_artifacts_and_preserves_hashes(tmp_path: Path):
    questionnaire = tmp_path / "filled.docx"
    make_questionnaire(questionnaire)
    materials = tmp_path / "materials"
    materials.mkdir()
    (materials / "数据结构课件.txt").write_text("chapter 1\nchapter 2", encoding="utf-8")
    Image.new("RGB", (32, 24), "white").save(materials / "实验样图.png")
    (materials / "damaged.pdf").write_bytes(b"bad-pdf")
    original_hashes = {path.name: sha256_file(path) for path in materials.iterdir()}

    data_root = tmp_path / "project_data"
    data_root.mkdir()
    database = Database(tmp_path / "workbench.sqlite")
    pipeline = WorkbenchPipeline(data_root, database, DisabledProvider(), ocr=StubOCR())
    package_id = pipeline.create_package(questionnaire, materials)
    product = pipeline.process(package_id)
    package = data_root / package_id
    try:
        for relative in (
            "reports/questionnaire_report.json", "reports/source_inventory.json",
            "reports/comparison_report.json", "reports/audit_log.jsonl", "reports/audit_report.md",
            "product.json", "buyer_card.json", "buyer_card.md",
            "catalog.md",
        ):
            assert (package / relative).exists(), relative
        copied_hashes = {path.name: sha256_file(path) for path in (package / "original" / "materials").iterdir()}
        assert copied_hashes == original_hashes
        assert product.quality_gate["approval_enabled"] is False
        assert database.get_package(package_id)["status"] == PackageStatus.NEEDS_SELLER_CONFIRMATION.value
        card_text = (package / "buyer_card.json").read_text(encoding="utf-8")
        assert "synthetic-contact" not in card_text
        assert "20元" not in card_text
    finally:
        cleanup_read_only(package)
